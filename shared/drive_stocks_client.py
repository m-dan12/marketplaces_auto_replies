"""Клиент для получения остатков товаров из Google Drive (xlsx файлы)."""

import io
from pathlib import Path
from typing import Dict

import openpyxl
from google.auth.transport.requests import Request
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload


class DriveStocksClient:
    """Клиент для загрузки остатков из Google Drive."""

    SCOPES = ['https://www.googleapis.com/auth/drive.readonly']
    TARGET_FILE = "Остатки.xlsx"

    def __init__(self, service_account_path: Path, target_folder: str, grandparent_folder: str = "Ozon"):
        """
        Инициализация клиента.

        Args:
            service_account_path: Путь к JSON файлу сервисного аккаунта
            target_folder: Название папки на Google Drive с файлом остатков
                (например, "Кабинет 3 (Milky Garden)")
            grandparent_folder: Название родительской папки для target_folder
                ("Ozon" или "Wildberries") — папки кабинетов называются одинаково
                под обоими маркетплейсами, поэтому без этого второго уровня поиск
                по имени неоднозначен.
        """
        self.credentials = service_account.Credentials.from_service_account_file(
            str(service_account_path), scopes=self.SCOPES
        )
        self.service = build('drive', 'v3', credentials=self.credentials)
        self.TARGET_FOLDER = target_folder
        self.GRANDPARENT_FOLDER = grandparent_folder

    def _find_file(self) -> str | None:
        """
        Поиск файла по имени и цепочке родительских папок (кабинет -> маркетплейс).

        Returns:
            ID файла или None если не найден
        """
        # Сначала ищем все файлы с нужным именем
        query = f"name='{self.TARGET_FILE}' and trashed=false"
        results = self.service.files().list(
            q=query,
            fields='files(id, name, parents)'
        ).execute()

        files = results.get('files', [])
        if not files:
            print(f"Файл '{self.TARGET_FILE}' не найден")
            return None

        # Проверяем каждый файл: совпадает ли папка кабинета и её родитель-маркетплейс
        for file_info in files:
            parent_id = file_info.get('parents', [None])[0]
            if not parent_id:
                continue

            parent_info = self.service.files().get(
                fileId=parent_id,
                fields='name, parents'
            ).execute()

            if parent_info.get('name') != self.TARGET_FOLDER:
                continue

            grandparent_id = parent_info.get('parents', [None])[0]
            if grandparent_id:
                grandparent_info = self.service.files().get(
                    fileId=grandparent_id, fields='name'
                ).execute()
                if grandparent_info.get('name') != self.GRANDPARENT_FOLDER:
                    continue

            print(f"Найден файл в папке '{self.GRANDPARENT_FOLDER}/{self.TARGET_FOLDER}'")
            return file_info['id']

        print(f"Файл '{self.TARGET_FILE}' не найден в папке '{self.GRANDPARENT_FOLDER}/{self.TARGET_FOLDER}'")
        return None

    def _download_file(self, file_id: str) -> bytes:
        """
        Загрузка файла из Google Drive.

        Args:
            file_id: ID файла в Google Drive

        Returns:
            Содержимое файла в байтах
        """
        request = self.service.files().get_media(fileId=file_id)
        file_buffer = io.BytesIO()
        downloader = MediaIoBaseDownload(file_buffer, request)

        done = False
        while not done:
            status, done = downloader.next_chunk()
            if status:
                print(f"Загрузка: {int(status.progress() * 100)}%")

        file_buffer.seek(0)
        return file_buffer.read()

    def _parse_ozon_stocks_xlsx(self, file_content: bytes) -> Dict[str, tuple[int, int]]:
        """
        Парсинг xlsx файла с остатками Ozon.

        Группирует FBO/FBS записи для каждого offer_id, суммируя остатки.

        Args:
            file_content: Содержимое xlsx файла

        Returns:
            Словарь {offer_id: (sku, total_stock)}
        """
        workbook = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
        sheet = workbook.active

        # Временное хранилище для группировки
        grouped = {}  # {offer_id: {'sku': int, 'stock': int}}
        headers = None

        for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
            # Первая строка - заголовки
            if i == 1:
                headers = row
                # Ищем индексы нужных колонок.
                # Колонка F "sku" — реальный артикул Ozon для ссылки/поиска товара;
                # "product_id" оставлен как legacy-фолбэк для файлов, где sku ещё не добавлена.
                try:
                    offer_id_idx = headers.index('offer_id')
                    stock_idx = headers.index('present')
                except ValueError as e:
                    print(f"Не найдены нужные колонки в заголовках: {headers}")
                    raise ValueError("Ожидались колонки: offer_id, sku, present") from e

                if 'sku' in headers:
                    sku_idx = headers.index('sku')
                elif 'product_id' in headers:
                    sku_idx = headers.index('product_id')
                    print("⚠️  Колонка 'sku' не найдена, используем устаревшую 'product_id'")
                else:
                    raise ValueError("Не найдена колонка 'sku' (или устаревшая 'product_id')")
                continue

            # Пропускаем пустые строки
            if not row or not row[offer_id_idx]:
                continue

            offer_id = str(row[offer_id_idx]).strip()
            sku_value = row[sku_idx]
            stock_value = row[stock_idx]

            # Парсим sku
            try:
                sku = int(float(sku_value)) if sku_value else 0
            except (ValueError, TypeError):
                print(f"Не удалось распарсить sku: {sku_value}")
                continue

            # Парсим остаток
            try:
                stock = int(float(stock_value)) if stock_value else 0
            except (ValueError, TypeError):
                stock = 0

            # Группируем: суммируем FBO + FBS
            if offer_id not in grouped:
                grouped[offer_id] = {'sku': sku, 'stock': stock}
            else:
                grouped[offer_id]['stock'] += stock

        workbook.close()

        # Преобразуем в финальный формат
        result = {
            offer_id: (data['sku'], data['stock'])
            for offer_id, data in grouped.items()
        }

        return result

    def _parse_wb_stocks_xlsx(self, file_content: bytes) -> Dict[str, tuple[int, int]]:
        """
        Парсинг xlsx файла с остатками Wildberries.

        Формат отличается от Ozon: колонки `nmId`, `vendorCode` и суммарный
        остаток по всем складам в колонке "Всего находится на складах"
        (выгрузка из личного кабинета WB, а не наш собственный формат).

        Args:
            file_content: Содержимое xlsx файла

        Returns:
            Словарь {vendor_code: (nm_id, total_stock)}
        """
        workbook = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
        sheet = workbook.active

        result: Dict[str, tuple[int, int]] = {}
        nm_id_idx = vendor_code_idx = stock_idx = None

        for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
            if i == 1:
                headers = row
                try:
                    nm_id_idx = headers.index('nmId')
                    vendor_code_idx = headers.index('vendorCode')
                    stock_idx = headers.index('Всего находится на складах')
                except ValueError as e:
                    print(f"Не найдены нужные колонки в заголовках: {headers}")
                    raise ValueError(
                        "Ожидались колонки: nmId, vendorCode, Всего находится на складах"
                    ) from e
                continue

            if not row or not row[vendor_code_idx]:
                continue

            vendor_code = str(row[vendor_code_idx]).strip()
            nm_id_value = row[nm_id_idx]
            stock_value = row[stock_idx]

            try:
                nm_id = int(float(nm_id_value)) if nm_id_value else 0
            except (ValueError, TypeError):
                print(f"Не удалось распарсить nmId: {nm_id_value}")
                continue

            try:
                stock = int(float(stock_value)) if stock_value else 0
            except (ValueError, TypeError):
                stock = 0

            result[vendor_code] = (nm_id, stock)

        workbook.close()
        return result

    def get_stocks(self) -> Dict[str, tuple[int, int]]:
        """
        Получение остатков Ozon из файла на Google Drive.

        Returns:
            Словарь {offer_id: (sku, total_stock)}
        """
        print(f"Поиск файла '{self.TARGET_FILE}' в папке '{self.TARGET_FOLDER}'")
        file_id = self._find_file()

        if not file_id:
            raise FileNotFoundError(
                f"Файл '{self.TARGET_FILE}' не найден в папке '{self.TARGET_FOLDER}'"
            )

        print(f"Загрузка файла (ID: {file_id})...")
        file_content = self._download_file(file_id)

        print("Парсинг остатков из xlsx...")
        stocks = self._parse_ozon_stocks_xlsx(file_content)

        print(f"Загружено остатков для {len(stocks)} товаров")
        return stocks

    def get_wb_stocks(self) -> Dict[str, tuple[int, int]]:
        """
        Получение остатков Wildberries из файла на Google Drive.

        Returns:
            Словарь {vendor_code: (nm_id, total_stock)}
        """
        print(f"Поиск файла '{self.TARGET_FILE}' в папке '{self.TARGET_FOLDER}'")
        file_id = self._find_file()

        if not file_id:
            raise FileNotFoundError(
                f"Файл '{self.TARGET_FILE}' не найден в папке '{self.TARGET_FOLDER}'"
            )

        print(f"Загрузка файла (ID: {file_id})...")
        file_content = self._download_file(file_id)

        print("Парсинг остатков из xlsx...")
        stocks = self._parse_wb_stocks_xlsx(file_content)

        print(f"Загружено остатков для {len(stocks)} товаров")
        return stocks


def load_stocks_from_drive(drive_folder: str) -> Dict[str, tuple[int, int]]:
    """
    Загрузка остатков Ozon из Google Drive для указанного кабинета.

    Args:
        drive_folder: Название папки на Google Drive с файлом остатков
            (например, "Кабинет 3 (Milky Garden)")

    Returns:
        Словарь {offer_id: (sku, total_stock)}
    """
    client = DriveStocksClient(_service_account_path(), target_folder=drive_folder, grandparent_folder="Ozon")
    return client.get_stocks()


def load_wb_stocks_from_drive(drive_folder: str) -> Dict[str, tuple[int, int]]:
    """
    Загрузка остатков Wildberries из Google Drive для указанного кабинета.

    Args:
        drive_folder: Название папки на Google Drive с файлом остатков
            (например, "Кабинет 2 (Milky Garden)" — нумерация кабинетов
            на WB отличается от Ozon, см. `wb_drive_folder` в `config/accounts.py`)

    Returns:
        Словарь {vendor_code: (nm_id, total_stock)}
    """
    client = DriveStocksClient(_service_account_path(), target_folder=drive_folder, grandparent_folder="Wildberries")
    return client.get_wb_stocks()


def _service_account_path() -> Path:
    """Путь к сервисному аккаунту Google Drive (общий для Ozon и Wildberries)."""
    path = Path(__file__).parent.parent / "drive_service_account.json"

    if not path.exists():
        raise FileNotFoundError(f"Файл сервисного аккаунта не найден: {path}")

    return path
